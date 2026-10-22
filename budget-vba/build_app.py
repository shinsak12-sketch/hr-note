# -*- coding: utf-8 -*-
"""사업비 관리 프로그램 (VBA 테스트판) 생성: 프로그램 xlsm + 샘플 데이터 + zip."""
import sys, io, os, zipfile, random, math, datetime as dt
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import Rule, DataBarRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.comments import Comment
from openpyxl.worksheet.properties import Outline

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from vbaproj import build_vba_project

OUTDIR = sys.argv[1]
TEST = "--test" in sys.argv          # 자동 점검용: 메시지 창 없이 동작

# ───────── 디자인 토큰 ─────────
FONT = "맑은 고딕"
INK, INK2, MUTE = "1E2A26", "4B5A55", "8A9692"
PRI, PRI2, PRI3 = "1B4D3E", "2E6E5A", "E6F0EB"
LINE, LINE2 = "E1E8E4", "9DB8AD"
IN_FILL, IN_FONT = "FFFBEA", "1F4FD1"
SUB_FILL, KEY_FILL = "EDF4F0", "F4F9F6"
DIS_FILL, DIS_FONT = "F1F3F2", "B7BFBB"
RED, GREEN, AMBER = "B42318", "1A7F37", "8A5A00"
MONEY_K = '#,##0,;△#,##0,;"-"'
MONEY_W = '#,##0;△#,##0;"-"'
MONEY_M = '#,##0.0,,;△#,##0.0,,;"-"'
PCT = '0.0%;△0.0%;"-"'
MONTH_FMT = '0"월"'

def F(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)
def side(c, style="thin"):
    return Side(style=style, color=c)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
LEFTW = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
LEFT0 = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
TB = Border(bottom=side(LINE))
BOX = Border(left=side(LINE2), right=side(LINE2), top=side(LINE2), bottom=side(LINE2))
UNLOCK = Protection(locked=False)

def put(ws, ref, v=None, font=None, bg=None, al=None, fmt=None, bd=None, unlock=False):
    c = ws[ref]
    if v is not None:
        c.value = v
    c.font = font or F()
    if bg: c.fill = fill(bg)
    if al: c.alignment = al
    if fmt: c.number_format = fmt
    if bd: c.border = bd
    if unlock: c.protection = UNLOCK
    return c

def dxf(fill_=None, color=None, bold=None, numfmt=None, nid=None, border=None):
    kw = {}
    if fill_: kw["fill"] = fill(fill_)
    if color or bold:
        kw["font"] = Font(color=color, bold=bold)
    if numfmt: kw["numFmt"] = NumberFormat(numFmtId=nid, formatCode=numfmt)
    if border: kw["border"] = border
    return DifferentialStyle(**kw)

def cf(ws, rng, formula, stop=False, **kw):
    ws.conditional_formatting.add(rng, Rule(type="expression", formula=[formula], stopIfTrue=stop, dxf=dxf(**kw)))

def unit_rules(ws, rng):
    cf(ws, rng, 'DispUnit="원"', numfmt=MONEY_W, nid=401)
    cf(ws, rng, 'DispUnit="백만원"', numfmt=MONEY_M, nid=402)

def base_sheet(ws, tab, widths):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False
    ws.sheet_view.zoomScale = 100
    pass  # 탭은 숨기므로 탭 색 없음
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    ws.protection.sheet = not TEST
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.autoFilter = False

def nav_row(ws, last_col):
    ws.row_dimensions[1].height = 36
    for c in range(1, last_col + 1):
        ws.cell(row=1, column=c).fill = fill(PRI)
    ws.row_dimensions[2].height = 8

def title(ws, ref, text, sub_ref=None, sub=None):
    put(ws, ref, text, F(20, True, PRI), al=LEFT0)
    ws.row_dimensions[ws[ref].row].height = 36
    if sub_ref:
        put(ws, sub_ref, sub, F(10, color=MUTE), al=LEFT0)
        ws.row_dimensions[ws[sub_ref].row].height = 18

def section(ws, ref, text, through=None):
    put(ws, ref, text, F(11.5, True, PRI), al=LEFT0, bd=Border(bottom=side(PRI, "medium")))
    if through:
        r = ws[ref].row
        for c in range(ws[ref].column + 1, ws[through].column + 1):
            ws.cell(row=r, column=c).border = Border(bottom=side(PRI, "medium"))
    ws.row_dimensions[ws[ref].row].height = 24

# ───────── 샘플 계정·계획 (2026년) ─────────
Z12 = [None] * 12
ACCOUNTS = [
    # 코드, 대, 중, 소, 전년(2025) 실적, 방식, 증감률, 금액·단가, 인원·수량, 개월, 분할, 분할값, 비고
    ("1101", "인건비", "급여", "기본급", 1_140_000_000, "인원", None, 4_300_000, 23, 12, "균등", Z12, "23명(신규 1명 포함)"),
    ("1102", "인건비", "급여", "상여금", 236_000_000, "%증감", 0.03, None, None, None, "수기",
        [None]*5 + [121_540_000] + [None]*5 + [121_540_000], "6·12월 지급"),
    ("1103", "인건비", "급여", "제수당", 118_000_000, "%증감", 0.03, None, None, None, "균등", Z12, ""),
    ("1104", "인건비", "퇴직급여", "퇴직급여", 102_000_000, "%증감", 0.05, None, None, None, "균등", Z12, ""),
    ("1105", "인건비", "4대보험", "4대보험", 131_000_000, "%증감", 0.04, None, None, None, "균등", Z12, "보험요율 인상 반영"),
    ("1106", "인건비", "복리후생", "복리후생비", 39_600_000, "인원", None, 150_000, 23, 12, "균등", Z12, ""),
    ("2101", "지급수수료", "외주", "조사 위탁수수료", 412_000_000, "%증감", 0.08, None, None, None, "균등", Z12, "위탁 물량 증가"),
    ("2102", "지급수수료", "자문", "의료자문료", 68_400_000, "단가×수량", None, 150_000, 40, 12, "균등", Z12, "월 40건"),
    ("2103", "지급수수료", "자문", "법률자문료", 48_000_000, "절대증감", None, 6_000_000, None, None, "균등", Z12, "소송 1건 추가 예상"),
    ("2104", "지급수수료", "전산", "전산 사용료", 150_000_000, "단가×수량", None, 13_000_000, 1, 12, "균등", Z12, ""),
    ("3101", "여비·차량", "여비교통", "현장출장비", 92_000_000, "%증감", 0.06, None, None, None, "균등", Z12, ""),
    ("3102", "여비·차량", "차량", "차량유지비", 46_800_000, "단가×수량", None, 450_000, 9, 12, "균등", Z12, "업무차량 9대"),
    ("3103", "여비·차량", "차량", "유류비", 58_000_000, "%증감", 0.02, None, None, None, "균등", Z12, ""),
    ("4101", "임차·관리", "임차", "사무실 임차료", 300_000_000, "단가×수량", None, 25_000_000, 1, 12, "균등", Z12, ""),
    ("4102", "임차·관리", "관리", "관리비", 81_600_000, "절대증감", None, 2_400_000, None, None, "균등", Z12, ""),
    ("5101", "일반관리", "통신", "통신비", 22_800_000, "%증감", 0.0, None, None, None, "균등", Z12, ""),
    ("5102", "일반관리", "사무", "소모품비", 28_400_000, "%증감", -0.05, None, None, None, "균등", Z12, "전자결재 확대로 절감"),
    ("5103", "일반관리", "교육", "교육훈련비", 30_500_000, "수기", None, 36_000_000, None, None, "비율",
        [None, None, 0.15, 0.10, 0.10, 0.05, None, None, 0.15, 0.15, 0.20, 0.10], "상·하반기 집중"),
    ("5104", "일반관리", "사무", "도서인쇄비", 6_200_000, "%증감", 0.0, None, None, None, "균등", Z12, ""),
]
METHODS = ["%증감", "절대증감", "인원", "단가×수량", "수기"]
SPLITS = ["균등", "수기", "비율"]

def xround(x, unit=1000):
    return int((Decimal(x) / unit).quantize(Decimal(1), rounding=ROUND_HALF_UP)) * unit

def calc_plan(a):
    code, l1, l2, l3, prev, meth, pct, amt, qty, mon, split, sv, note = a
    if meth == "%증감": k = prev * (1 + pct)
    elif meth == "절대증감": k = prev + amt
    elif meth in ("인원", "단가×수량"): k = amt * qty * (mon or 1)
    else: k = amt
    months = []
    for i in range(12):
        if split == "수기": raw = sv[i] or 0
        elif split == "비율": raw = k * (sv[i] or 0)
        else: raw = k / 12
        months.append(xround(raw))
    return k, months

# ───────── 프로그램 통합 문서 ─────────
wb = Workbook()
wb.code_name = "ThisWorkbook"
try:
    wb._named_styles["Normal"].font = Font(name=FONT, size=10)
except Exception:
    pass
home = wb.active; home.title = "홈"
plan = wb.create_sheet("사업비계획")
inp = wb.create_sheet("실적입력")
comp = wb.create_sheet("계획대비실적")
led = wb.create_sheet("실적원장")
acct = wb.create_sheet("계정관리")
sys_ = wb.create_sheet("SYS")
logs = wb.create_sheet("LOG")
for ws, cn in [(home, "shHome"), (plan, "shPlan"), (inp, "shInput"), (comp, "shComp"), (led, "shLedger"),
               (acct, "shAcct"), (sys_, "shSys"), (logs, "shLog")]:
    ws.sheet_properties.codeName = cn
sys_.sheet_state = "hidden"
logs.sheet_state = "hidden"
wb.views[0].showSheetTabs = False

def add_name(name, ref):
    dn = DefinedName(name, attr_text=ref)
    try:
        wb.defined_names[name] = dn
    except TypeError:
        wb.defined_names.append(dn)

add_name("PlanYear", "'SYS'!$B$2")
add_name("DispUnit", "'홈'!$C$19")
add_name("RndLabel", "'홈'!$C$20")
add_name("RndMode", "'홈'!$C$21")
add_name("BaseMonth", "'계획대비실적'!$C$5")
add_name("UnitDiv", "'SYS'!$B$10")
add_name("RndUnit", "'SYS'!$B$11")
add_name("UnitFmt", "'SYS'!$B$12")
add_name("AcctList", "OFFSET('SYS'!$H$2,0,0,MAX(1,COUNTA('SYS'!$H$2:$H$600)),1)")
add_name("YearList", "OFFSET('SYS'!$J$2,0,0,MAX(1,COUNT('SYS'!$J$2:$J$101)),1)")
add_name("MonthList", "'SYS'!$L$2:$L$13")

def dv_list(ws, src, cells, err="목록에서 고르세요."):
    dv = DataValidation(type="list", formula1=src, allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = err, "입력 확인", True
    ws.add_data_validation(dv)
    for c in cells:
        dv.add(c)

PLAN_FIRST, PLAN_LAST = 8, 600
COMP_FIRST, COMP_LAST = 17, 609
LED_FIRST, LED_LAST = 7, 20006
ACCT_FIRST, ACCT_LAST = 7, 306
P = "'사업비계획'"
LG = "'실적원장'"

# ═════════ SYS (숨김) ═════════
sys_.column_dimensions["A"].width = 18
sys_.column_dimensions["B"].width = 30
for ref, lab in [("A2", "열린 연도"), ("A3", "데이터 파일"), ("A4", "변경 있음"), ("A5", "마지막 저장"),
                 ("A6", "다음 실적 번호"), ("A7", "모드(1개발 2점검)"), ("A8", "마지막 메시지"), ("A9", "점검용 입력값"),
                 ("A10", "표시 단위 배수"), ("A11", "반올림 단위(원)"), ("A12", "문구용 숫자 서식"),
                 ("A13", "연간 계획"), ("A14", "누계 계획"), ("A15", "누계 집행"), ("A16", "연간 집행"), ("A17", "전년 실적")]:
    sys_[ref] = lab
sys_["B4"] = False
sys_["B6"] = 1
sys_["B7"] = 2 if TEST else 0
sys_["B9"] = 2027
sys_["B10"] = '=IF(DispUnit="원",1,IF(DispUnit="백만원",1000000,1000))'
sys_["B11"] = "=IFERROR(VLOOKUP(RndLabel,$F$2:$G$7,2,FALSE),1000)"
sys_["B12"] = '=IF(DispUnit="백만원","#,##0.0","#,##0")'
sys_["B13"] = f"=SUMIFS({P}!$L$8:$L$600,{P}!$AP$8:$AP$600,\"D\")"
sys_["B14"] = "=SUMIF('계획대비실적'!$C$7:$N$7,\"<=\"&BaseMonth,'계획대비실적'!$C$8:$N$8)"
sys_["B15"] = f"=SUMIFS({LG}!$G$7:$G$20006,{LG}!$C$7:$C$20006,\"<=\"&BaseMonth)"
sys_["B16"] = f"=SUM({LG}!$G$7:$G$20006)"
sys_["B17"] = f"=SUMIFS({P}!$E$8:$E$600,{P}!$AP$8:$AP$600,\"D\")"
for i, (lab, v) in enumerate([("원", 1), ("십원", 10), ("백원", 100), ("천원", 1000), ("만원", 10000), ("십만원", 100000)]):
    sys_[f"F{2+i}"] = lab
    sys_[f"G{2+i}"] = v
sys_["H1"] = "계정 목록"
sys_["J1"] = "연도 목록"
for i in range(12):
    sys_[f"L{2+i}"] = i + 1
for col in "BJ":
    pass
sys_["A700"] = "계획 수식 견본"
sys_["A702"] = "계획대비실적 수식 견본"

# 계획 수식 견본 (700행) — VBA가 FormulaR1C1 로 각 계정 줄에 복사. 목록(H2:H600 등)과 겹치지 않게 아래쪽에 둔다
def rnd(raw):
    return (f'IF(RndMode="절상",ROUNDUP(({raw})/RndUnit,0)*RndUnit,'
            f'IF(RndMode="절하",ROUNDDOWN(({raw})/RndUnit,0)*RndUnit,ROUND(({raw})/RndUnit,0)*RndUnit))')
SIN = [CL(15 + i) for i in range(12)]
MON = [CL(28 + i) for i in range(12)]
r = 700
E, Fm, G, H, I, J, N, K = (f"${c}{r}" for c in ("E", "F", "G", "H", "I", "J", "N", "K"))
sys_[f"K{r}"] = (f'=IF({Fm}="%증감",{E}*(1+{G}),IF({Fm}="절대증감",{E}+{H},IF(OR({Fm}="인원",{Fm}="단가×수량"),'
                 f'{H}*{I}*IF({J}>0,{J},1),IF({Fm}="수기",{H},0))))')
sys_[f"L{r}"] = f"=SUM(AB{r}:AM{r})"
sys_[f"M{r}"] = f'=IF({E}=0,"",$L{r}/{E}-1)'
srng = f"$O{r}:$Z{r}"
sys_[f"AA{r}"] = (f'=IF({N}="비율",IF(ABS(SUM({srng})-1)<0.0005,"OK","합 "&TEXT(SUM({srng}),"0.0%")),'
                  f'IF({N}="수기",IF(ROUND(SUM({srng})-{K},0)=0,"OK","차 "&TEXT((SUM({srng})-{K})/UnitDiv,UnitFmt)),""))')
for i in range(12):
    s = f"${SIN[i]}{r}"
    sys_[f"{MON[i]}{r}"] = "=" + rnd(f'IF({N}="수기",{s},IF({N}="비율",{K}*{s},{K}/12))')
sys_[f"AN{r}"] = ("=" +
    f'IF({Fm}="%증감","전년 "&TEXT({E}/UnitDiv,UnitFmt)&" × (1"&IF({G}<0,"−","+")&TEXT(ABS({G}),"0.0%")&")",'
    f'IF({Fm}="절대증감","전년 "&TEXT({E}/UnitDiv,UnitFmt)&IF({H}<0," − "," + ")&TEXT(ABS({H})/UnitDiv,UnitFmt),'
    f'IF({Fm}="인원","인당 "&TEXT({H}/UnitDiv,UnitFmt)&" × "&{I}&"명"&IF({J}>1," × "&{J}&"개월",""),'
    f'IF({Fm}="단가×수량","단가 "&TEXT({H}/UnitDiv,UnitFmt)&" × "&{I}&IF({J}>1," × "&{J}&"개월",""),'
    f'IF({Fm}="수기","연간 "&TEXT({H}/UnitDiv,UnitFmt),"")))))'
    f'&IF({N}="수기"," · 월 수기분할",IF({N}="비율"," · 월 비율분할",""))')

# 계획대비실적 수식 견본 (702행)
r = 702
A_, Q_ = f"$A{r}", f"$Q{r}"
sys_[f"Q{r}"] = f"=IFERROR(MATCH({A_},{P}!$A:$A,0),0)"
sys_[f"D{r}"] = f"=IF({Q_}=0,0,INDEX({P}!$L:$L,{Q_}))"
sys_[f"E{r}"] = f"=IF({Q_}=0,0,INDEX({P}!$AB:$AM,{Q_},BaseMonth))"
sys_[f"F{r}"] = f"=SUMIFS({LG}!$G:$G,{LG}!$E:$E,{A_},{LG}!$C:$C,BaseMonth)"
sys_[f"G{r}"] = f"=$F{r}-$E{r}"
sys_[f"H{r}"] = f'=IF($E{r}=0,"",$F{r}/$E{r})'
sys_[f"I{r}"] = f"=IF({Q_}=0,0,SUM(INDEX({P}!$AB:$AM,{Q_},1):INDEX({P}!$AB:$AM,{Q_},BaseMonth)))"
sys_[f"J{r}"] = f'=SUMIFS({LG}!$G:$G,{LG}!$E:$E,{A_},{LG}!$C:$C,"<="&BaseMonth)'
sys_[f"K{r}"] = f"=$J{r}-$I{r}"
sys_[f"L{r}"] = f'=IF($I{r}=0,"",$J{r}/$I{r})'
sys_[f"M{r}"] = f'=IF($D{r}=0,"",$J{r}/$D{r})'
sys_[f"N{r}"] = f"=$D{r}-$J{r}"
sys_[f"O{r}"] = f'=IF(AND($D{r}=0,$J{r}=0),"",IF($J{r}>$D{r},"연간 초과",IF($J{r}>$I{r},"계획 초과","정상")))'

# LOG
logs["A1"], logs["B1"], logs["C1"] = "시각", "사용자", "내용"

# ═════════ 홈 ═════════
base_sheet(home, PRI, {"A": 2, "B": 15, "C": 20, "D": 13, "E": 4, "F": 21, "G": 21, "H": 21, "I": 21, "J": 2})
nav_row(home, 10)
title(home, "B3", "사업비 관리 프로그램", "B4",
      '=IF(PlanYear="","열린 데이터가 없습니다.  [새 연도 만들기]로 시작하세요.",'
      'PlanYear&"년 데이터가 열려 있습니다.   기준월 "&BaseMonth&"월   ·   금액 단위 "&DispUnit)')
put(home, "B5", "⚠ 지금은 매크로가 꺼져 있습니다. 엑셀 위쪽 노란 줄의 [콘텐츠 사용]을 눌러야 버튼이 작동합니다.",
    F(10, True, "9A3412"), al=LEFT0)
home.row_dimensions[5].height = 20
section(home, "B7", "데이터", "D7")
rows = [(8, "열린 연도", '=IF(PlanYear="","-",PlanYear&"년")', F(13, True, PRI)),
        (9, "저장 상태", '=IF(PlanYear="","-",IF(SYS!$B$4,"저장 안 된 변경 있음","저장됨"))', F(10, True, INK)),
        (10, "마지막 저장", '=IF(SYS!$B$5="","-",SYS!$B$5)', F(10, color=INK2))]
for rr, lab, fml, fnt in rows:
    put(home, f"B{rr}", lab, F(10, True, INK2), al=LEFT, bd=TB)
    home.merge_cells(f"C{rr}:D{rr}")
    put(home, f"C{rr}", fml, fnt, al=LEFT, bd=TB)
    home[f"D{rr}"].border = TB
    home.row_dimensions[rr].height = 24
cf(home, "C9", 'ISNUMBER(SEARCH("안 된",$C$9))', color="9A3412", bold=True, fill_="FDEBD8")
cf(home, "C9", '$C$9="저장됨"', color=GREEN, bold=True)
put(home, "B12", "연도 선택", F(10, True, INK2), al=LEFT)
put(home, "C12", None, F(12, True, IN_FONT), IN_FILL, CENTER, '0"년"', BOX, unlock=True)
home.row_dimensions[12].height = 28
home.row_dimensions[13].height = 8
home.row_dimensions[14].height = 30
home.row_dimensions[15].height = 6
home.row_dimensions[16].height = 30
dv_list(home, "YearList", ["C12"])
section(home, "B18", "설정  (연도별로 저장)", "D18")
for rr, lab, val, lst, hint in [(19, "금액 표시 단위", "천원", '"원,천원,백만원"', "입력은 항상 원"),
                                (20, "반올림 단위", "천원", '"원,십원,백원,천원,만원,십만원"', "월별 계획 금액"),
                                (21, "반올림 방식", "반올림", '"반올림,절상,절하"', "확정 = 월 합계")]:
    put(home, f"B{rr}", lab, F(10, True, INK2), al=LEFT, bd=TB)
    put(home, f"C{rr}", val, F(11, True, IN_FONT), IN_FILL, CENTER, bd=BOX, unlock=True)
    put(home, f"D{rr}", hint, F(9, color=MUTE), al=LEFT, bd=TB)
    home.row_dimensions[rr].height = 26
    dv_list(home, lst, [f"C{rr}"])
home.merge_cells("B23:D24")
put(home, "B23", '="데이터 파일: "&IF(SYS!$B$3="","-",SYS!$B$3)', F(8.5, color=MUTE), al=Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1))

section(home, "F7", "한눈에 보기", "I7")
put(home, "I7", '=IF(PlanYear="","","기준월 "&BaseMonth&"월 · "&DispUnit)', F(9, color=MUTE), al=RIGHT,
    bd=Border(bottom=side(PRI, "medium")))
TILES = [("F", "연간 계획", "=SYS!$B$13", MONEY_K, '="전년 실적 "&TEXT(SYS!$B$17/UnitDiv,UnitFmt)'),
         ("G", "누계 집행", "=SYS!$B$15", MONEY_K, '="1~"&BaseMonth&"월 계획 "&TEXT(SYS!$B$14/UnitDiv,UnitFmt)'),
         ("H", "누계 집행률", '=IF(SYS!$B$14=0,"",SYS!$B$15/SYS!$B$14)', PCT,
          '=IF(SYS!$B$14=0,"",IF(SYS!$B$15>SYS!$B$14,"계획보다 "&TEXT((SYS!$B$15-SYS!$B$14)/UnitDiv,UnitFmt)&" 더 씀","계획보다 "&TEXT((SYS!$B$14-SYS!$B$15)/UnitDiv,UnitFmt)&" 덜 씀"))'),
         ("I", "연간 잔액", "=SYS!$B$13-SYS!$B$16", MONEY_K, '=IF(SYS!$B$13=0,"","연간 진도 "&TEXT(SYS!$B$16/SYS!$B$13,"0.0%"))')]
for col, lab, fml, fmt, sub in TILES:
    for rr in range(8, 12):
        c = home[f"{col}{rr}"]
        c.fill = fill(KEY_FILL)
        c.border = Border(left=side("FFFFFF", "thick"), right=side("FFFFFF", "thick"),
                          top=side(PRI, "medium") if rr == 8 else None)
    put(home, f"{col}8", lab, F(9.5, True, INK2), KEY_FILL, LEFT, bd=Border(top=side(PRI, "medium"), left=side("FFFFFF", "thick"), right=side("FFFFFF", "thick")))
    home.merge_cells(f"{col}9:{col}10")
    put(home, f"{col}9", fml, F(18, True, PRI), KEY_FILL, LEFT, fmt, Border(left=side("FFFFFF", "thick"), right=side("FFFFFF", "thick")))
    put(home, f"{col}11", sub, F(8.5, color=MUTE), KEY_FILL, LEFT, bd=Border(left=side("FFFFFF", "thick"), right=side("FFFFFF", "thick")))
for rr, h in [(8, 20), (9, 18), (10, 18), (11, 18)]:
    home.row_dimensions[rr].height = max(home.row_dimensions[rr].height or 0, h)
unit_rules(home, "F9:G10")
unit_rules(home, "I9:I10")
cf(home, "H9:H10", "AND(ISNUMBER($H$9),$H$9>1)", color=RED)
cf(home, "I9:I10", "$I$9<0", color=RED)
section(home, "F13", "메뉴", "I13")
for rr in range(14, 23):
    home.row_dimensions[rr].height = max(home.row_dimensions[rr].height or 0, 22)

# ═════════ 사업비계획 ═════════
PW = {1: 7, 2: 10, 3: 9, 4: 14.5, 5: 12.5, 6: 10, 7: 7.5, 8: 11.5, 9: 7.5, 10: 5.5, 11: 12.5,
      12: 13.5, 13: 7.5, 14: 6.5, 27: 9, 40: 36, 41: 22, 42: 4}
for i in range(12):
    PW[15 + i] = 8.5
    PW[28 + i] = 10.2
base_sheet(plan, PRI, {CL(k): v for k, v in PW.items()})
plan.sheet_view.zoomScale = 90
nav_row(plan, 42)
title(plan, "A3", '=IF(PlanYear="","사업비 계획",PlanYear&"년 사업비 계획")', "A4",
      '="금액 단위: "&DispUnit&"   ·   입력은 원 단위   ·   월별 금액은 "&RndLabel&" 단위 "&RndMode&" 후 합계가 확정 계획액   ·   계정 추가·삭제는 [계정 관리]에서"')
plan.row_dimensions[2].height = 8
put(plan, "A5", "범례", F(9, True, MUTE), al=LEFT0)
put(plan, "B5", "입력", F(9, True, IN_FONT), IN_FILL, CENTER)
put(plan, "C5", "자동계산", F(9, True, INK), "FFFFFF", CENTER, bd=BOX)
put(plan, "D5", "해당 없음", F(9, True, DIS_FONT), DIS_FILL, CENTER)
put(plan, "E5", "  [+] / [−] 로 '산출 입력'·'월 분할 입력' 열을 펴고 접습니다.", F(9, color=MUTE), al=LEFT0)
plan.row_dimensions[5].height = 20
bands = [("A", "D", "계정", "DCE8E2"), ("E", "E", "전년", "E7EAE9"), ("F", "K", "산출 입력", "FFF0C2"),
         ("L", "N", "계획", "CFE6DA"), ("O", "AA", "월 분할 입력  (수기 = 금액 · 비율 = %)", "FCE6C9"),
         ("AB", "AM", "월별 확정 계획 (자동)", "DCEDE4"), ("AN", "AO", "근거 · 메모", "E7EAE9")]
for a, b, t, c in bands:
    if a != b:
        plan.merge_cells(f"{a}6:{b}6")
    put(plan, f"{a}6", t, F(9.5, True, PRI), c, CENTER)
    for ci in range(plan[f"{a}6"].column, plan[f"{b}6"].column + 1):
        plan.cell(row=6, column=ci).fill = fill(c)
plan.row_dimensions[6].height = 20
heads = {1: "코드", 2: "대분류", 3: "중분류", 4: "소분류", 5: "전년 실적", 6: "산출 방식", 7: "증감률", 8: "금액·단가",
         9: "인원·수량", 10: "개월", 11: "산출 금액", 12: "확정\n계획액", 13: "전년비", 14: "분할", 27: "분할 점검",
         40: "산출근거 (자동)", 41: "비고"}
for k, t in heads.items():
    put(plan, f"{CL(k)}7", t, F(9.5, True, "FFFFFF"), PRI, CENTER)
for i in range(12):
    put(plan, f"{SIN[i]}7", f"{i+1}월", F(9, True, "FFFFFF"), "8A6A2E", CENTER)
    put(plan, f"{MON[i]}7", f"{i+1}월", F(9.5, True, "FFFFFF"), PRI2, CENTER)
plan.row_dimensions[7].height = 34
for k, txt in {6: "%증감 · 절대증감 · 인원 · 단가×수량 · 수기 중 선택.\n쓰지 않는 입력칸은 회색으로 바뀝니다.",
               7: "%증감에서만 사용. 3% 처럼 % 를 붙여 입력.",
               8: "원 단위.\n절대증감 = 증감액(감액은 −)\n인원 = 인당 월액\n단가×수량 = 단가\n수기 = 연간 금액",
               9: "인원 = 인원 수 / 단가×수량 = 수량", 10: "인원·단가×수량에서 곱할 개월 수 (비우면 1)",
               14: "균등 = 산출 금액 ÷ 12\n수기 = 월별 금액 직접 입력\n비율 = 월별 % 입력",
               12: "월별 확정 계획(반올림 후)의 합계", 27: "비율 합 100%, 수기 합 = 산출 금액인지 확인"}.items():
    cm = Comment(txt, "사업비 관리")
    cm.width, cm.height = 250, 110
    plan[f"{CL(k)}7"].comment = cm

INPUT_COLS = set(range(5, 11)) | {14} | set(range(15, 27)) | {41}
for rr in range(PLAN_FIRST, PLAN_LAST + 1):
    for c in range(1, 43):
        cell = plan.cell(row=rr, column=c)
        cell.border = TB
        if c in INPUT_COLS:
            cell.fill = fill(IN_FILL)
            cell.font = F(10, color=IN_FONT)
        elif c == 12:
            cell.fill = fill(KEY_FILL)
            cell.font = F(10.5, True)
        elif c in (11, 27, 40):
            cell.font = F(9.5 if c != 11 else 10, color=INK2)
        else:
            cell.font = F(10)
        if c in (5, 8, 11, 12) or 15 <= c <= 26 or 28 <= c <= 39:
            cell.number_format = MONEY_K
            cell.alignment = RIGHT
        elif c in (7, 13):
            cell.number_format = PCT
            cell.alignment = RIGHT
        elif c in (9,):
            cell.number_format = "#,##0"
            cell.alignment = RIGHT
        elif c == 10:
            cell.number_format = "0"
            cell.alignment = RIGHT
        elif c in (1, 6, 14, 27):
            cell.alignment = CENTER
            if c == 1:
                cell.number_format = "@"
        elif c == 40:
            cell.alignment = LEFT
        else:
            cell.alignment = LEFT0
    plan.row_dimensions[rr].height = 21
    if rr > PLAN_FIRST:
        plan.row_dimensions[rr].hidden = True
plan.freeze_panes = "E8"
plan.sheet_properties.outlinePr = Outline(summaryBelow=True, summaryRight=True)
for c in range(6, 12):
    plan.column_dimensions[CL(c)].outline_level = 1
for c in range(15, 28):
    plan.column_dimensions[CL(c)].outline_level = 1
    plan.column_dimensions[CL(c)].hidden = True
plan.column_dimensions["AP"].hidden = True
dv_list(plan, '"' + ",".join(METHODS) + '"', [f"F{PLAN_FIRST}:F{PLAN_LAST}"])
dv_list(plan, '"' + ",".join(SPLITS) + '"', [f"N{PLAN_FIRST}:N{PLAN_LAST}"])

R0 = PLAN_FIRST
AREA = f"A{R0}:AO{PLAN_LAST}"
SPL = f"O{R0}:Z{PLAN_LAST}"
cf(plan, SPL, f'AND($AP{R0}="D",$N{R0}="비율")', stop=True, numfmt=PCT, nid=403)
for rng in [f"E{R0}:E{PLAN_LAST}", f"H{R0}:H{PLAN_LAST}", f"K{R0}:L{PLAN_LAST}", SPL, f"AB{R0}:AM{PLAN_LAST}"]:
    unit_rules(plan, rng)
cf(plan, AREA, f'$AP{R0}="S"', stop=True, fill_=SUB_FILL, color=PRI, bold=True,
   border=Border(top=side(LINE2), bottom=side(LINE2)))
cf(plan, AREA, f'$AP{R0}="T"', stop=True, fill_=PRI, color="FFFFFF", bold=True)
def dis(col, cond):
    cf(plan, f"{col}{R0}:{col}{PLAN_LAST}", f'AND($AP{R0}="D",{cond})', fill_=DIS_FILL, color=DIS_FONT)
dis("G", f'$F{R0}<>"%증감"')
dis("H", f'OR($F{R0}="%증감",$F{R0}="")')
dis("I", f'AND($F{R0}<>"인원",$F{R0}<>"단가×수량")')
dis("J", f'AND($F{R0}<>"인원",$F{R0}<>"단가×수량")')
cf(plan, SPL, f'AND($AP{R0}="D",$N{R0}<>"수기",$N{R0}<>"비율")', fill_=DIS_FILL, color=DIS_FONT)
cf(plan, f"N{R0}:N{PLAN_LAST}", f'AND($AP{R0}="D",$AA{R0}<>"",$AA{R0}<>"OK")', stop=True, fill_="FDE2E1", color=RED, bold=True)
for val, fc, bgc in [("균등", "1E6B4A", "E2F2E9"), ("수기", "5B3FA0", "EEE7FA"), ("비율", AMBER, "FCEFD4")]:
    cf(plan, f"N{R0}:N{PLAN_LAST}", f'$N{R0}="{val}"', fill_=bgc, color=fc, bold=True)
cf(plan, f"AA{R0}:AA{PLAN_LAST}", f'$AA{R0}="OK"', color=GREEN, bold=True)
cf(plan, f"AA{R0}:AA{PLAN_LAST}", f'AND($AA{R0}<>"",$AA{R0}<>"OK")', color=RED, bold=True)
cf(plan, f"M{R0}:M{PLAN_LAST}", f'AND(ISNUMBER($M{R0}),$M{R0}>0.0005)', color=RED)
cf(plan, f"M{R0}:M{PLAN_LAST}", f'AND(ISNUMBER($M{R0}),$M{R0}<-0.0005)', color=GREEN)
plan.page_setup.orientation = "landscape"
plan.page_setup.paperSize = plan.PAPERSIZE_A3
plan.page_setup.fitToWidth, plan.page_setup.fitToHeight = 1, 0
plan.sheet_properties.pageSetUpPr.fitToPage = True
plan.print_title_rows = "6:7"

# ═════════ 실적입력 ═════════
base_sheet(inp, "8A6A2E", {"A": 2, "B": 13, "C": 30, "D": 15, "E": 15, "F": 3, "G": 17, "H": 15, "I": 15, "J": 15, "K": 2, "L": 10})
nav_row(inp, 12)
inp.column_dimensions["L"].hidden = True
title(inp, "B3", "실적 입력", "B4",
      '=IF(PlanYear="","열린 데이터가 없습니다.",PlanYear&"년   ·   금액은 원 단위로 입력   ·   표시 단위 "&DispUnit&"   ·   등록하면 [실적 원장]에 한 줄씩 쌓입니다")')
section(inp, "B5", "새 실적", "E5")
FORM = [(6, "월", None, MONTH_FMT, "MonthList"), (7, "일자", None, "yyyy-mm-dd", None),
        (8, "계정", None, None, "AcctList"), (9, "금액 (원)", None, "#,##0", None),
        (10, "거래처", None, None, None), (11, "비고", None, None, None)]
for rr, lab, v, fmt, lst in FORM:
    put(inp, f"B{rr}", lab, F(10, True, INK2), al=LEFT)
    put(inp, f"C{rr}", v, F(11, rr in (6, 8, 9), IN_FONT), IN_FILL, LEFT if rr != 9 else RIGHT, fmt, BOX, unlock=True)
    inp.row_dimensions[rr].height = 26
    if lst:
        dv_list(inp, lst, [f"C{rr}"])
put(inp, "D7", "선택 (비워도 됨)", F(9, color=MUTE), al=LEFT)
inp.merge_cells("D9:E9")
put(inp, "D9", '=IF(N(C9)=0,"","=  "&TEXT(C9/UnitDiv,UnitFmt)&" "&DispUnit)', F(11, True, PRI), al=LEFT)
inp.row_dimensions[12].height = 34
inp.merge_cells("B14:E14")
put(inp, "B14", None, F(10, True, GREEN), al=LEFT)
inp.row_dimensions[13].height = 6
inp.row_dimensions[14].height = 22
# 선택 계정 현황
inp.merge_cells("G5:J5")
put(inp, "G5", '=IF($L$6="","선택 계정 현황","선택 계정   ·   "&$C$8)', F(11.5, True, PRI), al=LEFT0,
    bd=Border(bottom=side(PRI, "medium")))
for c in "HIJ":
    inp[f"{c}5"].border = Border(bottom=side(PRI, "medium"))
for c, t in zip("GHIJ", ["", "계획", "집행", "남은 금액"]):
    put(inp, f"{c}6", t, F(9.5, True, "FFFFFF"), PRI2, CENTER)
inp["L6"] = '=IF($C$8="","",LEFT($C$8,FIND(" ",$C$8&" ")-1))'
inp["L7"] = f"=IFERROR(MATCH($L$6,{P}!$A:$A,0),0)"
inp["L16"] = f"=COUNT({LG}!$B$7:$B$20006)"
SUMA = f"SUMIFS({LG}!$G:$G,{LG}!$E:$E,$L$6"
STAT = [(7, '=IF(ISNUMBER($C$6),$C$6&"월","해당 월")',
         f"=IF(OR($L$7=0,NOT(ISNUMBER($C$6))),0,INDEX({P}!$AB:$AM,$L$7,$C$6))",
         f'=IF(OR($L$6="",NOT(ISNUMBER($C$6))),0,{SUMA},{LG}!$C:$C,$C$6))'),
        (8, '=IF(ISNUMBER($C$6),"누계 1~"&$C$6&"월","누계")',
         f"=IF(OR($L$7=0,NOT(ISNUMBER($C$6))),0,SUM(INDEX({P}!$AB:$AM,$L$7,1):INDEX({P}!$AB:$AM,$L$7,$C$6)))",
         f'=IF(OR($L$6="",NOT(ISNUMBER($C$6))),0,{SUMA},{LG}!$C:$C,"<="&$C$6))'),
        (9, "연간", f"=IF($L$7=0,0,INDEX({P}!$L:$L,$L$7))", f'=IF($L$6="",0,{SUMA}))'),
        (10, "등록하면 (월)", "=H7", "=I7+N($C$9)")]
for rr, lab, h, i_ in STAT:
    put(inp, f"G{rr}", lab, F(10, True, INK2), KEY_FILL, LEFT, bd=TB)
    put(inp, f"H{rr}", h, F(10.5, rr == 9), None, RIGHT, MONEY_K, TB)
    put(inp, f"I{rr}", i_, F(10.5, rr == 9), None, RIGHT, MONEY_K, TB)
    put(inp, f"J{rr}", f"=H{rr}-I{rr}", F(10.5, True), None, RIGHT, MONEY_K, TB)
put(inp, "G10", "등록하면 (월)", F(10, True, PRI), "FFF4D1", LEFT, bd=TB)
for c in "HIJ":
    inp[f"{c}10"].fill = fill("FFF4D1")
unit_rules(inp, "H7:J10")
cf(inp, "J7:J10", "J7<0", color=RED, bold=True)
inp.merge_cells("G11:J12")
put(inp, "G11", '=IF(OR($L$6="",N($C$9)=0),"",IF(I10>H10,"⚠  등록하면 "&$C$6&"월 계획을 "&TEXT((I10-H10)/UnitDiv,UnitFmt)&" "&DispUnit&" 넘습니다.","✓  "&$C$6&"월 계획 안입니다.  남는 금액 "&TEXT(J10/UnitDiv,UnitFmt)&" "&DispUnit))',
    F(10.5, True, INK), al=LEFTW)
cf(inp, "G11:J12", 'LEFT($G$11,1)="⚠"', fill_="FDE2E1", color=RED, bold=True)
cf(inp, "G11:J12", 'LEFT($G$11,1)="✓"', fill_="E2F2E9", color=GREEN, bold=True)
# 최근 입력
section(inp, "B15", "최근 입력 (최신 15건)", "J15")
inp.row_dimensions[15].height = 30
for c, t in zip("BCDEFGH", ["번호", "계정", "월", "금액", "", "거래처", "비고"]):
    put(inp, f"{c}16", t, F(9.5, True, "FFFFFF"), PRI2, CENTER)
inp.merge_cells("H16:J16")
IDX = "$L$16-(ROW()-17)"
for rr in range(17, 32):
    put(inp, f"B{rr}", f'=IF({IDX}<1,"",INDEX({LG}!$B$7:$B$20006,{IDX}))', F(10, color=MUTE), None, CENTER, "0", TB)
    put(inp, f"C{rr}", f'=IF($B{rr}="","",INDEX({LG}!$E$7:$E$20006,{IDX})&" "&INDEX({LG}!$F$7:$F$20006,{IDX}))', F(10), None, LEFT, None, TB)
    put(inp, f"D{rr}", f'=IF($B{rr}="","",INDEX({LG}!$C$7:$C$20006,{IDX}))', F(10), None, CENTER, MONTH_FMT, TB)
    put(inp, f"E{rr}", f'=IF($B{rr}="","",INDEX({LG}!$G$7:$G$20006,{IDX}))', F(10, True), None, RIGHT, MONEY_K, TB)
    put(inp, f"F{rr}", None, F(10), None, None, None, TB)
    put(inp, f"G{rr}", f'=IF($B{rr}="","",INDEX({LG}!$H$7:$H$20006,{IDX})&"")', F(10), None, LEFT, None, TB)
    inp.merge_cells(f"H{rr}:J{rr}")
    put(inp, f"H{rr}", f'=IF($B{rr}="","",INDEX({LG}!$I$7:$I$20006,{IDX})&"")', F(10, color=INK2), None, LEFT, None, TB)
    for c in "IJ":
        inp[f"{c}{rr}"].border = TB
    inp.row_dimensions[rr].height = 21
unit_rules(inp, "E17:E31")

# ═════════ 계획대비실적 ═════════
CW = {"A": 7, "B": 11, "C": 16, "D": 13, "E": 12, "F": 12, "G": 12, "H": 9, "I": 13, "J": 13, "K": 12,
      "L": 9, "M": 14, "N": 13, "O": 11, "P": 4, "Q": 4}
base_sheet(comp, PRI2, CW)
nav_row(comp, 17)
comp.column_dimensions["P"].hidden = True
comp.column_dimensions["Q"].hidden = True
title(comp, "A3", '=IF(PlanYear="","계획 대비 실적",PlanYear&"년 계획 대비 실적")', "A4",
      '="기준월 "&BaseMonth&"월   ·   금액 단위 "&DispUnit&"   ·   집행 = [실적 원장] 합계   ·   차이 = 집행 − 계획 (＋면 더 씀)"')
put(comp, "B5", "기준월", F(10, True, INK2), al=LEFT)
put(comp, "C5", 9, F(12, True, IN_FONT), IN_FILL, CENTER, MONTH_FMT, BOX, unlock=True)
put(comp, "D5", "← 바꾸면 아래 표가 바로 다시 계산됩니다", F(9, color=MUTE), al=LEFT0)
comp.row_dimensions[5].height = 26
dv_list(comp, "MonthList", ["C5"])
put(comp, "B7", "월별", F(9.5, True, "FFFFFF"), PRI, CENTER)
MCOLS = [CL(3 + i) for i in range(12)]
for i, c in enumerate(MCOLS):
    put(comp, f"{c}7", i + 1, F(9.5, True, "FFFFFF"), PRI, CENTER, MONTH_FMT)
put(comp, "O7", "누계", F(9.5, True, "FFFFFF"), PRI, CENTER)
labels = ["계획", "집행", "차이", "집행률", "누계 집행률"]
for k, lab in enumerate(labels):
    rr = 8 + k
    put(comp, f"B{rr}", lab, F(10, True, INK2), KEY_FILL, LEFT, bd=TB)
    comp.row_dimensions[rr].height = 22
for i, c in enumerate(MCOLS):
    pm = MON[i]
    put(comp, f"{c}8", f'=SUMIFS({P}!{pm}$8:{pm}$600,{P}!$AP$8:$AP$600,"D")', F(10), None, RIGHT, MONEY_K, TB)
    put(comp, f"{c}9", f"=SUMIFS({LG}!$G$7:$G$20006,{LG}!$C$7:$C$20006,{c}$7)", F(10, True), None, RIGHT, MONEY_K, TB)
    put(comp, f"{c}10", f'=IF({c}$7>BaseMonth,"",{c}9-{c}8)', F(10), None, RIGHT, MONEY_K, TB)
    put(comp, f"{c}11", f'=IF({c}8=0,"",{c}9/{c}8)', F(10), None, RIGHT, PCT, TB)
    put(comp, f"{c}12", f'=IF({c}$7>BaseMonth,"",IF(SUM($C8:{c}8)=0,"",SUM($C9:{c}9)/SUM($C8:{c}8)))', F(10, color=INK2), None, RIGHT, PCT, TB)
put(comp, "O8", '=SUMIF($C$7:$N$7,"<="&BaseMonth,C8:N8)', F(10, True), KEY_FILL, RIGHT, MONEY_K, TB)
put(comp, "O9", '=SUMIF($C$7:$N$7,"<="&BaseMonth,C9:N9)', F(10, True), KEY_FILL, RIGHT, MONEY_K, TB)
put(comp, "O10", "=O9-O8", F(10, True), KEY_FILL, RIGHT, MONEY_K, TB)
put(comp, "O11", '=IF(O8=0,"",O9/O8)', F(10, True), KEY_FILL, RIGHT, PCT, TB)
put(comp, "O12", None, F(10), KEY_FILL, None, None, TB)
unit_rules(comp, "C8:O10")
cf(comp, "C7:N7", "C$7=BaseMonth", fill_="C98500", color="FFFFFF", bold=True)
cf(comp, "C8:N12", "C$7=BaseMonth", fill_="FFF4D1")
cf(comp, "C9:N12", "C$7>BaseMonth", color=DIS_FONT)
cf(comp, "C10:O10", "AND(ISNUMBER(C10),C10>0)", color=RED, bold=True)
cf(comp, "C10:O10", "AND(ISNUMBER(C10),C10<0)", color=GREEN)
cf(comp, "C11:O12", "AND(ISNUMBER(C11),C11>1)", color=RED, bold=True)

section(comp, "A14", "계정별", "O14")
put(comp, "D14", "상태:  정상 = 누계 계획 안  ·  계획 초과 = 누계 계획보다 더 씀  ·  연간 초과 = 연간 계획도 넘음",
    F(9, color=MUTE), al=LEFT0, bd=Border(bottom=side(PRI, "medium")))
cbands = [("A", "C", "계정", "DCE8E2"), ("D", "D", "연간", "E7EAE9"), ("E", "H", '=BaseMonth&"월 (당월)"', "FFF0C2"),
          ("I", "L", '="누계 (1~"&BaseMonth&"월)"', "CFE6DA"), ("M", "N", "연간 진도", "E7EAE9"), ("O", "O", "상태", "E7EAE9")]
for a, b, t, c in cbands:
    if a != b:
        comp.merge_cells(f"{a}15:{b}15")
    put(comp, f"{a}15", t, F(9.5, True, PRI), c, CENTER)
    for ci in range(comp[f"{a}15"].column, comp[f"{b}15"].column + 1):
        comp.cell(row=15, column=ci).fill = fill(c)
comp.row_dimensions[15].height = 20
for c, t in zip("ABCDEFGHIJKLMNO", ["코드", "대분류", "계정", "연간 계획", "계획", "집행", "차이", "집행률",
                                    "계획", "집행", "차이", "집행률", "집행률", "잔액", "상태"]):
    put(comp, f"{c}16", t, F(9.5, True, "FFFFFF"), PRI, CENTER)
comp.row_dimensions[16].height = 26
for rr in range(COMP_FIRST, COMP_LAST + 1):
    for ci in range(1, 16):
        cell = comp.cell(row=rr, column=ci)
        cell.border = TB
        cell.font = F(10, ci in (6, 10))
        if ci in (4, 5, 6, 7, 9, 10, 11, 14):
            cell.number_format = MONEY_K
            cell.alignment = RIGHT
        elif ci in (8, 12, 13):
            cell.number_format = PCT
            cell.alignment = RIGHT
        elif ci in (1, 15):
            cell.alignment = CENTER
            if ci == 1:
                cell.number_format = "@"
        else:
            cell.alignment = LEFT0
    comp.row_dimensions[rr].height = 21
    if rr > COMP_FIRST:
        comp.row_dimensions[rr].hidden = True
C0 = COMP_FIRST
CAREA = f"A{C0}:O{COMP_LAST}"
for rng in [f"D{C0}:G{COMP_LAST}", f"I{C0}:K{COMP_LAST}", f"N{C0}:N{COMP_LAST}"]:
    unit_rules(comp, rng)
cf(comp, CAREA, f'$P{C0}="S"', stop=True, fill_=SUB_FILL, color=PRI, bold=True,
   border=Border(top=side(LINE2), bottom=side(LINE2)))
cf(comp, CAREA, f'$P{C0}="T"', stop=True, fill_=PRI, color="FFFFFF", bold=True)
cf(comp, f"G{C0}:G{COMP_LAST}", f"AND(ISNUMBER(G{C0}),G{C0}>0)", color=RED, bold=True)
cf(comp, f"K{C0}:K{COMP_LAST}", f"AND(ISNUMBER(K{C0}),K{C0}>0)", color=RED, bold=True)
cf(comp, f"G{C0}:G{COMP_LAST}", f"AND(ISNUMBER(G{C0}),G{C0}<0)", color=GREEN)
cf(comp, f"K{C0}:K{COMP_LAST}", f"AND(ISNUMBER(K{C0}),K{C0}<0)", color=GREEN)
cf(comp, f"H{C0}:H{COMP_LAST}", f"AND(ISNUMBER(H{C0}),H{C0}>1)", color=RED, bold=True)
cf(comp, f"L{C0}:L{COMP_LAST}", f"AND(ISNUMBER(L{C0}),L{C0}>1)", color=RED, bold=True)
cf(comp, f"N{C0}:N{COMP_LAST}", f"AND(ISNUMBER(N{C0}),N{C0}<0)", color=RED, bold=True)
cf(comp, f"O{C0}:O{COMP_LAST}", f'$O{C0}="연간 초과"', fill_="FDE2E1", color=RED, bold=True)
cf(comp, f"O{C0}:O{COMP_LAST}", f'$O{C0}="계획 초과"', fill_="FCEFD4", color=AMBER, bold=True)
cf(comp, f"O{C0}:O{COMP_LAST}", f'$O{C0}="정상"', fill_="E2F2E9", color=GREEN, bold=True)
comp.conditional_formatting.add(f"M{C0}:M{COMP_LAST}",
    DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="7FB8A2"))
comp.freeze_panes = "D17"
comp.page_setup.orientation = "landscape"
comp.page_setup.paperSize = comp.PAPERSIZE_A4
comp.page_setup.fitToWidth, comp.page_setup.fitToHeight = 1, 0
comp.sheet_properties.pageSetUpPr.fitToPage = True

# ═════════ 실적원장 ═════════
LW = {"A": 2, "B": 8, "C": 7, "D": 11, "E": 7, "F": 16, "G": 14, "H": 18, "I": 34, "J": 7, "K": 16, "L": 10}
base_sheet(led, "8A6A2E", LW)
nav_row(led, 12)
title(led, "B3", '=IF(PlanYear="","실적 원장",PlanYear&"년 실적 원장")', "B4",
      "모든 실적이 한 줄씩 쌓입니다.  월 · 일자 · 금액 · 거래처 · 비고는 여기서 바로 고칠 수 있습니다.  머리글의 ▼로 걸러 볼 수 있습니다.")
led.merge_cells("B5:I5")
put(led, "B5", '="보이는 줄  "&TEXT(SUBTOTAL(2,B7:B20006),"#,##0")&"건   ·   합계 "&TEXT(SUBTOTAL(9,G7:G20006)/UnitDiv,UnitFmt)&" "&DispUnit',
    F(10.5, True, PRI), al=LEFT0)
led.row_dimensions[5].height = 30
for c, t in zip("BCDEFGHIJKL", ["번호", "월", "일자", "코드", "계정", "금액", "거래처", "비고", "출처", "입력 시각", "입력자"]):
    put(led, f"{c}6", t, F(9.5, True, "FFFFFF"), PRI, CENTER)
led.row_dimensions[6].height = 24
colstyle = {"B": ("0", CENTER, F(10, color=MUTE), False), "C": (MONTH_FMT, CENTER, F(10), True),
            "D": ("yyyy-mm-dd", CENTER, F(10), True), "E": ("@", CENTER, F(10), False),
            "F": ("General", LEFT0, F(10), False), "G": (MONEY_K, RIGHT, F(10, True), True),
            "H": ("General", LEFT, F(10), True), "I": ("General", LEFT, F(10, color=INK2), True),
            "J": ("General", CENTER, F(9, color=MUTE), False), "K": ("yyyy-mm-dd hh:mm", CENTER, F(9, color=MUTE), False),
            "L": ("General", CENTER, F(9, color=MUTE), False)}
for c, (fmt, al, fnt, editable) in colstyle.items():
    cd = led.column_dimensions[c]
    cd.number_format = fmt
    cd.alignment = al
    cd.font = fnt
    if editable:
        cd.protection = UNLOCK
cf(led, f"B{LED_FIRST}:L{LED_LAST}", f'$B{LED_FIRST}<>""', border=Border(bottom=side(LINE)))
unit_rules(led, f"G{LED_FIRST}:G{LED_LAST}")
led.auto_filter.ref = f"B6:L{LED_LAST}"
led.freeze_panes = "A7"

# ═════════ 계정관리 ═════════
base_sheet(acct, "8A9692", {"A": 2, "B": 10, "C": 14, "D": 14, "E": 20, "F": 40, "G": 2})
nav_row(acct, 7)
title(acct, "B3", "계정 관리", "B4", "계정을 추가 · 수정 · 삭제한 뒤 [계획에 반영]을 누르세요.  코드 순서대로 정렬되고, 대분류가 같은 계정끼리 소계가 묶입니다.")
acct.row_dimensions[5].height = 34
for c, t in zip("BCDEF", ["코드", "대분류", "중분류", "소분류", "메모"]):
    put(acct, f"{c}6", t, F(9.5, True, "FFFFFF"), PRI, CENTER)
acct.row_dimensions[6].height = 24
for rr in range(ACCT_FIRST, ACCT_LAST + 1):
    for c in "BCDEF":
        put(acct, f"{c}{rr}", None, F(10, c == "B", IN_FONT), IN_FILL, CENTER if c == "B" else LEFT0,
            "@" if c == "B" else None, TB, unlock=True)
    acct.row_dimensions[rr].height = 20
cf(acct, f"B{ACCT_FIRST}:B{ACCT_LAST}", f'AND($B{ACCT_FIRST}<>"",COUNTIF($B${ACCT_FIRST}:$B${ACCT_LAST},$B{ACCT_FIRST})>1)',
   fill_="FDE2E1", color=RED, bold=True)
acct.freeze_panes = "A7"

for w in (home, inp, led, acct):
    w.page_setup.orientation = "landscape"
    w.page_setup.paperSize = w.PAPERSIZE_A4
    w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
buf = io.BytesIO()
wb.save(buf)

# ───────── 도형 (상단 메뉴 + 화면 버튼) ─────────
EMU_PX, EMU_PT = 9525, 12700
def px(v):
    return int(round(v * EMU_PX))

def body(paras, anchor="ctr", align="ctr", lins=8):
    ps = []
    for text, size, bold, color in paras:
        ps.append(f'<a:p><a:pPr algn="{align}"/><a:r><a:rPr lang="ko-KR" altLang="en-US" sz="{int(size*100)}" b="{1 if bold else 0}">'
                  f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:latin typeface="{FONT}"/><a:ea typeface="{FONT}"/>'
                  f'</a:rPr><a:t>{text}</a:t></a:r></a:p>')
    return (f'<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" wrap="square" lIns="{px(lins)}" rIns="{px(lins)}" '
            f'tIns="0" bIns="0" rtlCol="0" anchor="{anchor}"/><a:lstStyle/>' + "".join(ps) + "</xdr:txBody>")

def sp(sid, name, macro, fill_c, paras, geom="roundRect", adj=16000, line=None, anchor="ctr", align="ctr", lins=8):
    m = f' macro="[0]!{macro}"' if macro else ""
    fl = f'<a:solidFill><a:srgbClr val="{fill_c}"/></a:solidFill>' if fill_c else "<a:noFill/>"
    ln = f'<a:ln w="9525"><a:solidFill><a:srgbClr val="{line}"/></a:solidFill></a:ln>' if line else "<a:ln><a:noFill/></a:ln>"
    gd = f'<a:gd name="adj" fmla="val {adj}"/>' if geom == "roundRect" else ""
    return (f'<xdr:sp{m} textlink=""><xdr:nvSpPr><xdr:cNvPr id="{sid}" name="{name}"/><xdr:cNvSpPr/></xdr:nvSpPr>'
            f'<xdr:spPr><a:prstGeom prst="{geom}"><a:avLst>{gd}</a:avLst></a:prstGeom>{fl}{ln}</xdr:spPr>'
            + body(paras, anchor, align, lins) + "</xdr:sp>")

def absolute(x, y, w, h, inner):
    return (f'<xdr:absoluteAnchor><xdr:pos x="{px(x)}" y="{px(y)}"/><xdr:ext cx="{px(w)}" cy="{px(h)}"/>'
            + inner + "<xdr:clientData/></xdr:absoluteAnchor>")

def in_cells(c1, r1, c2, r2, inner, widths, heights, m=4):
    """c1..c2 열(1부터), r1..r2 행 영역 안에 여백 m px 을 두고 배치"""
    wpx = int(widths.get(c2, 8.43) * 7 + 5)
    hpt = heights.get(r2, 15)
    return (f'<xdr:twoCellAnchor editAs="oneCell"><xdr:from><xdr:col>{c1-1}</xdr:col><xdr:colOff>{px(m)}</xdr:colOff>'
            f'<xdr:row>{r1-1}</xdr:row><xdr:rowOff>{px(m)}</xdr:rowOff></xdr:from>'
            f'<xdr:to><xdr:col>{c2-1}</xdr:col><xdr:colOff>{px(max(wpx - m, 1))}</xdr:colOff>'
            f'<xdr:row>{r2-1}</xdr:row><xdr:rowOff>{int(hpt*EMU_PT) - px(m)}</xdr:rowOff></xdr:to>'
            + inner + "<xdr:clientData/></xdr:twoCellAnchor>")

NAV = [("홈", "GoHome", 60), ("사업비 계획", "GoPlan", 100), ("실적 입력", "GoInput", 90),
       ("계획 대비 실적", "GoComp", 120), ("실적 원장", "GoLedger", 90), ("계정 관리", "GoAcct", 90)]
def nav_shapes(active):
    out = []
    sid = 2
    out.append(absolute(16, 0, 140, 48, sp(sid, "appname", None, None,
               [("사업비 관리", 13, True, "FFFFFF")], geom="rect", align="l", lins=0)))
    x = 160
    for label, macro, w in NAV:
        sid += 1
        on = (macro == active)
        out.append(absolute(x, 9, w, 30, sp(sid, f"nav_{macro}", macro, "FFFFFF" if on else PRI2,
                   [(label, 10, True, PRI if on else "E6F0EB")])))
        x += w + 6
    sid += 1
    out.append(absolute(x + 18, 9, 76, 30, sp(sid, "nav_save", "SaveData", "F2C14E", [("저장", 10, True, PRI)])))
    return out, sid

def col_widths(ws):
    out = {}
    for k, d in ws.column_dimensions.items():
        try:
            from openpyxl.utils import column_index_from_string
            out[column_index_from_string(k)] = d.width or 8.43
        except Exception:
            pass
    return out

def row_heights(ws):
    return {k: (d.height or 15) for k, d in ws.row_dimensions.items()}

def btn(sid, name, macro, label, c1, r1, c2, r2, ws, kind="pri", size=10.5):
    colors = {"pri": (PRI, "FFFFFF", None), "sec": ("FFFFFF", PRI, LINE2), "warn": ("FFFFFF", RED, "F1B7B2"),
              "save": ("F2C14E", PRI, None)}
    fc, tc, ln = colors[kind]
    return in_cells(c1, r1, c2, r2, sp(sid, name, macro, fc, [(label, size, True, tc)], line=ln),
                    col_widths(ws), row_heights(ws))

def tile(sid, name, macro, t1, t2, c1, r1, c2, r2, ws, fc):
    return in_cells(c1, r1, c2, r2, sp(sid, name, macro, fc, [(t1, 13, True, "FFFFFF"), (t2, 9.5, False, "D9E8E1")],
                    adj=8000, align="l", lins=16), col_widths(ws), row_heights(ws), m=5)

SHAPES = {}
s, sid = nav_shapes("GoHome")
s += [btn(sid + 1, "b_open", "OpenSelectedYear", "열기", 4, 12, 4, 12, home),
      btn(sid + 2, "b_new", "NewYear", "＋  새 연도 만들기", 2, 14, 4, 14, home, "sec"),
      btn(sid + 3, "b_save", "SaveData", "저장  (Ctrl+S)", 2, 16, 4, 16, home, "save"),
      tile(sid + 4, "t_plan", "GoPlan", "사업비 계획", "계정별 산출 방식 · 월 분할 · 확정 계획액", 6, 14, 7, 16, home, PRI),
      tile(sid + 5, "t_input", "GoInput", "실적 입력", "월 · 계정 · 금액을 넣고 [등록]", 8, 14, 9, 16, home, PRI),
      tile(sid + 6, "t_comp", "GoComp", "계획 대비 실적", "당월 · 누계 집행률, 더 쓴 계정 확인", 6, 17, 7, 19, home, PRI2),
      tile(sid + 7, "t_led", "GoLedger", "실적 원장", "등록한 실적 전체 · 고치기 · 지우기", 8, 17, 9, 19, home, PRI2),
      tile(sid + 8, "t_acct", "GoAcct", "계정 관리", "계정 추가 · 수정 후 계획에 반영", 6, 20, 7, 22, home, "5E6F69"),
      tile(sid + 9, "t_save", "SaveData", "저장", "data 폴더에 저장, 이전 파일은 backup", 8, 20, 9, 22, home, "8A6A2E")]
SHAPES["xl/worksheets/sheet1.xml"] = s
s, sid = nav_shapes("GoPlan")
s += [btn(sid + 1, "b_acct", "GoAcct", "계정 추가 · 수정  →", 40, 5, 40, 5, plan, "sec", 9.5)]
SHAPES["xl/worksheets/sheet2.xml"] = s
s, sid = nav_shapes("GoInput")
s += [btn(sid + 1, "b_reg", "RegisterActual", "등록", 3, 12, 3, 12, inp, "pri", 11.5),
      btn(sid + 2, "b_clr", "ClearInputForm", "지우기", 4, 12, 4, 12, inp, "sec"),
      btn(sid + 3, "b_del", "DeleteSelectedActual", "선택한 줄 삭제", 9, 15, 10, 15, inp, "warn", 9.5)]
SHAPES["xl/worksheets/sheet3.xml"] = s
s, sid = nav_shapes("GoComp")
SHAPES["xl/worksheets/sheet4.xml"] = s
s, sid = nav_shapes("GoLedger")
s += [btn(sid + 1, "b_del", "DeleteSelectedActual", "선택한 줄 삭제", 11, 5, 12, 5, led, "warn", 9.5)]
SHAPES["xl/worksheets/sheet5.xml"] = s
s, sid = nav_shapes("GoAcct")
s += [btn(sid + 1, "b_apply", "ApplyAccounts", "계획에 반영", 2, 5, 3, 5, acct, "pri", 11)]
SHAPES["xl/worksheets/sheet6.xml"] = s

if os.environ.get("NOSHAPES"):
    SHAPES = {}

def drawing_xml(parts):
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">' + "".join(parts) + "</xdr:wsDr>")

# ───────── VBA ─────────
VD = os.path.join(HERE, "vba")
def rd(n):
    return open(os.path.join(VD, n), encoding="utf-8").read()
DOC_ATTR = '''Attribute VB_Name = "{name}"
Attribute VB_Base = "0{{{guid}}}"
Attribute VB_GlobalNameSpace = False
Attribute VB_Creatable = False
Attribute VB_PredeclaredId = True
Attribute VB_Exposed = True
Attribute VB_TemplateDerived = False
Attribute VB_Customizable = True
'''
WB_GUID = "00020819-0000-0000-C000-000000000046"
WS_GUID = "00020820-0000-0000-C000-000000000046"
mods = [("ThisWorkbook", True, DOC_ATTR.format(name="ThisWorkbook", guid=WB_GUID) + ("" if os.environ.get("NODOC") else rd("ThisWorkbook.cls")))]
for cn in ["shHome", "shPlan", "shInput", "shComp", "shLedger", "shAcct", "shSys", "shLog"]:
    extra = rd(cn + ".cls") if os.path.exists(os.path.join(VD, cn + ".cls")) and not os.environ.get("NODOC") else ""
    mods.append((cn, True, DOC_ATTR.format(name=cn, guid=WS_GUID) + extra))
for m in os.environ.get("MODS", "App,DataIO,Screens").split(","):
    mods.append((m, False, rd(m + ".bas")))
vba_bin = build_vba_project(mods)

# ───────── xlsm 조립 ─────────
def assemble(xlsx_bytes, out_path):
    zin = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    names = zin.namelist()
    zout = zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED)
    ct_over = []
    for name in names:
        data = zin.read(name)
        if name in SHAPES:
            s = data.decode("utf-8")
            assert "<drawing " not in s, name
            idx = list(SHAPES).index(name) + 1
            tag = '<drawing r:id="rIdShp"/>'
            # drawing 은 legacyDrawing(메모) 앞, pageSetup 뒤
            if "<legacyDrawing" in s:
                s = s.replace("<legacyDrawing", tag + "<legacyDrawing", 1)
            elif "<tableParts" in s:
                s = s.replace("<tableParts", tag + "<tableParts", 1)
            else:
                s = s.replace("</worksheet>", tag + "</worksheet>")
            root_tag = s[s.index("<worksheet"):s.index(">", s.index("<worksheet"))]
            if "xmlns:r=" not in root_tag:
                s = s.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
            data = s.encode("utf-8")
        if name.endswith(".rels") and name.replace("_rels/", "").replace(".rels", "") in SHAPES:
            continue   # 아래에서 도형 관계를 더해 다시 쓴다
        zout.writestr(name, data)
    # 시트 관계 파일
    for i, name in enumerate(SHAPES, 1):
        rel = name.replace("worksheets/", "worksheets/_rels/") + ".rels"
        entry = (f'<Relationship Id="rIdShp" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" '
                 f'Target="../drawings/drawing{i}.xml"/>')
        if rel in names:
            s = zin.read(rel).decode("utf-8").replace("</Relationships>", entry + "</Relationships>")
        else:
            s = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                 '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + entry + "</Relationships>")
        zout.writestr(rel, s)
        zout.writestr(f"xl/drawings/drawing{i}.xml", drawing_xml(SHAPES[name]))
        ct_over.append(f'<Override PartName="/xl/drawings/drawing{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>')
    zout.writestr("xl/vbaProject.bin", vba_bin)
    zout.close()
    # Content types / workbook rels 수정 (다시 열어 고침)
    tmp = out_path + ".tmp"
    os.replace(out_path, tmp)
    zi = zipfile.ZipFile(tmp)
    zo = zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED)
    for item in zi.infolist():
        d = zi.read(item.filename)
        if item.filename == "[Content_Types].xml":
            s = d.decode("utf-8")
            s = s.replace("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                          "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            s = s.replace("<Default ", '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/><Default ', 1)
            s = s.replace("</Types>", "".join(ct_over) + "</Types>")
            d = s.encode("utf-8")
        elif item.filename == "xl/_rels/workbook.xml.rels":
            d = d.decode("utf-8").replace("</Relationships>",
                '<Relationship Id="rIdVBA" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
                'Target="vbaProject.bin"/></Relationships>').encode("utf-8")
        zo.writestr(item, d)
    zo.close()
    os.remove(tmp)

# ───────── 샘플 데이터 (2026년, 1~9월 실적) ─────────
def build_sample(path):
    random.seed(7)
    dw = Workbook()
    ws = dw.active
    ws.title = "정보"
    hdr_font = Font(name=FONT, bold=True, color="FFFFFF")
    def hdr(w, n):
        for c in range(1, n + 1):
            w.cell(row=1, column=c).font = hdr_font
            w.cell(row=1, column=c).fill = fill(PRI)
    plan_rows, ledger = [], []
    for a in ACCOUNTS:
        code, l1, l2, l3, prev, meth, pct, amt, qty, mon, split, sv, note = a
        k, months = calc_plan(a)
        row = [code, l1, l2, l3, prev, meth, pct, amt, qty, mon, k, sum(months), sum(months) / prev - 1, split] + \
              list(sv) + ["OK" if split in ("수기", "비율") else ""] + months + ["", note]
        plan_rows.append(row)
    # 실적 생성
    trend = {"3103": 1.12, "3101": 1.05, "5102": 0.82, "5101": 0.97, "2102": 1.03, "2101": 0.96, "1106": 1.04, "5104": 0.7}
    vendors = {"1101": "직원 급여", "1102": "직원 상여", "1103": "직원 수당", "1104": "퇴직연금 운용사", "1105": "4대보험 공단",
               "1106": "복지몰", "2101": "협력 조사업체", "2102": "자문의", "2103": "법무법인", "2104": "전산 운영사",
               "3101": "출장비 정산", "3102": "정비업체", "3103": "주유소(법인카드)", "4101": "건물주", "4102": "관리사무소",
               "5101": "통신사", "5102": "사무용품점", "5103": "교육기관", "5104": "서점"}
    multi = {"2102": 4, "3101": 3, "3103": 2, "5102": 2, "2101": 2}
    nid = 1
    for a, pr in zip(ACCOUNTS, plan_rows):
        code = a[0]
        months = pr[27:39]
        for m in range(1, 10):
            base = months[m - 1]
            if base == 0:
                continue
            f = trend.get(code, 1.0)
            if code in ("1101", "1102", "4101", "2104", "1104"):
                total = base
            else:
                total = base * f * random.uniform(0.94, 1.06)
            if code == "2103" and m == 4:
                total += 18_000_000
            parts = multi.get(code, 1)
            shares = [random.uniform(0.7, 1.3) for _ in range(parts)]
            ssum = sum(shares)
            for p_i, sh in enumerate(shares):
                amt = int(round(total * sh / ssum, -2))
                day = min(28, 5 + p_i * 7 + random.randint(0, 4))
                memo = ""
                if code == "2103" and m == 4 and p_i == 0:
                    memo = "소송 착수금 포함"
                elif code == "3103":
                    memo = "유가 상승"
                elif code == "1102":
                    memo = "상반기 상여"
                ts = dt.datetime(2026, m, day, 17, random.randint(0, 59))
                ledger.append([nid, m, dt.date(2026, m, day), code, a[3], amt, vendors[code], memo, "직접", ts, "경리담당"])
                nid += 1
    info = [("항목", "값"), ("연도", 2026), ("저장시각", "2026-10-15 18:02:11"), ("저장자", "경리담당"),
            ("프로그램", "사업비 관리 0.1 테스트"), ("표시단위", "천원"), ("반올림단위", "천원"), ("반올림방식", "반올림"),
            ("기준월", 9), ("다음ID", nid), ("계정 수", len(ACCOUNTS)), ("실적 건수", len(ledger))]
    for rr, (k, v) in enumerate(info, 1):
        ws.cell(row=rr, column=1, value=k)
        ws.cell(row=rr, column=2, value=v)
    hdr(ws, 2)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 26
    w2 = dw.create_sheet("계정")
    w2.append(["코드", "대분류", "중분류", "소분류", "메모"])
    for a in ACCOUNTS:
        w2.append([a[0], a[1], a[2], a[3], ""])
    for rr in range(2, len(ACCOUNTS) + 2):
        w2.cell(row=rr, column=1).number_format = "@"
    hdr(w2, 5)
    w3 = dw.create_sheet("계획")
    heads = ["코드", "대분류", "중분류", "소분류", "전년실적", "산출방식", "증감률", "금액단가", "인원수량", "개월", "산출금액",
             "확정계획액", "전년비", "분할"] + [f"분할{i}월" for i in range(1, 13)] + ["분할점검"] + \
            [f"계획{i}월" for i in range(1, 13)] + ["산출근거", "비고"]
    w3.append(heads)
    for row in plan_rows:
        w3.append(row)
    for rr in range(2, len(plan_rows) + 2):
        w3.cell(row=rr, column=1).number_format = "@"
        for c in [5, 8, 11, 12] + list(range(28, 40)):
            w3.cell(row=rr, column=c).number_format = "#,##0"
    hdr(w3, 41)
    w4 = dw.create_sheet("실적")
    w4.append(["ID", "월", "일자", "코드", "계정", "금액", "거래처", "비고", "출처", "입력시각", "입력자"])
    for row in ledger:
        w4.append(row)
    for rr in range(2, len(ledger) + 2):
        w4.cell(row=rr, column=4).number_format = "@"
        w4.cell(row=rr, column=6).number_format = "#,##0"
        w4.cell(row=rr, column=3).number_format = "yyyy-mm-dd"
        w4.cell(row=rr, column=10).number_format = "yyyy-mm-dd hh:mm"
    hdr(w4, 11)
    w5 = dw.create_sheet("이력")
    w5.append(["시각", "사용자", "내용"])
    w5.append(["2026-01-05 09:10:00", "경리담당", "2026년 새로 만듦 (샘플)"])
    w5.append(["2026-10-15 18:02:11", "경리담당", "저장"])
    hdr(w5, 3)
    dw.save(path)
    return len(ledger), plan_rows

os.makedirs(OUTDIR, exist_ok=True)
app_dir = os.path.join(OUTDIR, "BudgetApp")
os.makedirs(os.path.join(app_dir, "data"), exist_ok=True)
os.makedirs(os.path.join(app_dir, "backup"), exist_ok=True)
xlsm = os.path.join(app_dir, "BudgetManager.xlsm")
assemble(buf.getvalue(), xlsm)
nled, plan_rows = build_sample(os.path.join(app_dir, "data", "budget_2026.xlsx"))
print("saved", xlsm, "vba", len(vba_bin), "ledger", nled, "plan total", sum(r[11] for r in plan_rows))
